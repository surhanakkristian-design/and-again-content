#!/usr/bin/env python3
"""
apply_en.py — write the ENGLISH rows of a part from a compact content file.

    python3 apply_en.py <workbook.xlsx> <content.txt> [--check]

Content file (UTF-8, pipe-delimited; '#' lines and blank lines ignored):

    E <exercise_id>
    en|<intro_text>|<correct_answer>|<distractor_1>|<distractor_2>

Every exercise in the file is validated before anything is written, against the
workbook's own exercises sheet: id exists, type-specific intro rule (27/68/74/75
empty, 69 starts with "<word>" means, grammar types carry exactly one '...' gap),
English budget (intro <= 90, options <= 50), three distinct options, no four-dot gap,
no option that is a bare copy of another. The workbook is saved once. Only the
`en` row of sentence_translations is touched; nothing else is written.
"""
import re, sys
import openpyxl

NO_INTRO = {27, 68, 74, 75}
STEM = {69}
LIMITS = [("intro_text", 90), ("correct_answer", 50), ("distractor_1", 50), ("distractor_2", 50)]


def s(v):
    return "" if v is None else str(v).strip()


def canon(w):
    return re.sub(r"^(a|an|the|to)\s+", "", s(w).lower())


def main():
    wbp, cpath = sys.argv[1], sys.argv[2]
    check_only = "--check" in sys.argv
    errors, content = [], {}
    cur = None
    for n, raw in enumerate(open(cpath, encoding="utf-8"), 1):
        line = raw.rstrip("\n")
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if line.startswith("E "):
            cur = int(line[2:].strip())
            if cur in content:
                errors.append(f"line {n}: exercise {cur} listed twice")
            content[cur] = None
            continue
        parts = line.split("|")
        if len(parts) != 5 or parts[0].strip() != "en":
            errors.append(f"line {n}: expected 'en|intro|correct|d1|d2', got {len(parts)} fields")
            continue
        if cur is None:
            errors.append(f"line {n}: text row before any E header"); continue
        if content[cur] is not None:
            errors.append(f"line {n}: exercise {cur} has two en rows"); continue
        content[cur] = tuple(p.strip() for p in parts[1:])

    wb = openpyxl.load_workbook(wbp)
    ex = {}
    for r in wb["exercises"].iter_rows(min_row=2, values_only=True):
        if r[0] is None or not s(r[0]).isdigit():
            continue
        ex[int(r[0])] = (int(r[1]), int(r[2]), int(r[3]))
    con = {int(r[0]): s(r[1]) for r in wb["word_concepts"].iter_rows(min_row=2, values_only=True)
           if r[0] is not None and s(r[0]).isdigit()}
    # type 69 uses the media's `meaning of the word` (All Words, column F) verbatim as its
    # correct answer — the field is written to a 50-char budget for exactly this use.
    meaning = {int(r[0]): s(r[5]) for r in wb["All Words"].iter_rows(min_row=2, values_only=True)
               if r[0] is not None and s(r[0]).isdigit()}
    st = wb["sentence_translations"]
    en_row = {}
    for i, r in enumerate(st.iter_rows(min_row=2), start=2):
        if r[0].value is None or not s(r[0].value).isdigit():
            continue
        if s(r[2].value) == "en":
            en_row[int(r[1].value)] = i

    for eid, vals in content.items():
        if vals is None:
            errors.append(f"exercise {eid}: no en row"); continue
        if eid not in ex:
            errors.append(f"exercise {eid}: not in the workbook"); continue
        if eid not in en_row:
            errors.append(f"exercise {eid}: no en sentence_translations row"); continue
        cid, mid, t = ex[eid]
        intro, ca, d1, d2 = vals
        word = con.get(cid, "")
        for (name, lim), v in zip(LIMITS, vals):
            if len(v) > lim:
                errors.append(f"E13 exercise {eid}: en {name} {len(v)} chars > {lim}: {v!r}")
        if t in NO_INTRO and intro:
            errors.append(f"exercise {eid}: type {t} must have empty intro_text")
        if t in STEM:
            if not re.match(r'^"[^"]+" means\.\.\.$', intro):
                errors.append(f"exercise {eid}: type 69 stem must be '\"<word>\" means...' got {intro!r}")
            elif canon(intro.split('"')[1]) != word:
                errors.append(f"exercise {eid}: type 69 stem word {intro.split(chr(34))[1]!r} != concept {word!r}")
            if mid in meaning and ca != meaning[mid]:
                errors.append(f"exercise {eid}: type 69 correct_answer must be the media's meaning column "
                              f"verbatim — got {ca!r}, meaning is {meaning[mid]!r}")
        if t not in NO_INTRO and t not in STEM:
            if not intro:
                errors.append(f"exercise {eid}: type {t} needs intro_text")
            elif intro.count("...") != 1:
                errors.append(f"exercise {eid}: intro must contain exactly one '...' gap: {intro!r}")
        if "...." in intro:
            errors.append(f"exercise {eid}: four-dot gap marker")
        if not ca or not d1 or not d2:
            errors.append(f"exercise {eid}: all three options must be filled")
        if len({ca.lower(), d1.lower(), d2.lower()}) < 3:
            errors.append(f"exercise {eid}: options not distinct: {vals[1:]!r}")
        if t in (74, 75):
            for v in (ca, d1, d2):
                if len(v.split()) > 4:
                    errors.append(f"exercise {eid}: Simple Explanation option over 4 words: {v!r}")
            if word and word not in ca.lower() and canon(word) not in ca.lower():
                errors.append(f"exercise {eid}: Simple Explanation answer must contain {word!r}: {ca!r}")
        if t in (27, 68) and word and word not in ca.lower():
            errors.append(f"exercise {eid}: Label answer should contain {word!r}: {ca!r}")

    if errors:
        print(f"REFUSED — {len(errors)} problem(s), nothing written:")
        for e in errors[:80]:
            print("  " + e)
        sys.exit(1)
    if check_only:
        print(f"OK — {len(content)} exercise(s) pass, nothing written (--check)")
        return
    for eid, vals in content.items():
        row = en_row[eid]
        for off, v in enumerate(vals):
            st.cell(row=row, column=4 + off).value = v
    wb.save(wbp)
    print(f"OK — {len(content)} en row(s) written")


if __name__ == "__main__":
    main()
