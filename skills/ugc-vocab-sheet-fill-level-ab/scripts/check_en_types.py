#!/usr/bin/env python3
"""check_en_types.py — heuristic: does each English answer look like its exercise type?

    python3 check_en_types.py <workbook.xlsx>

Prints exercises whose correct_answer does not match a coarse pattern for the type
(e.g. Past Continuous without was/were + -ing, Question tags without '?'). A hit is a
row to eyeball, not an error. Read-only.
"""
import re, sys
import openpyxl

PAT = {
 1: r"\b(am|is|are|was|were|'m|'s|'re|be|been|being)\b", 2: r"\b(got|have|has|haven't|hasn't)\b",
 4: r"\b(can|can't|cannot)\b", 6: r"^(a|an|the|-)\b", 7: r"\b(this|that|these|those)\b",
 8: r"\b(my|your|his|her|its|our|their)\b", 9: r"^(who|what|where|when|why|how|which|whose)\b",
 10: r"^(in|on|at|under|next to|behind|in front of|between|near|from|to|by|over|above|below|after|before|until|during)\b",
 11: r"\b(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|thirteen|fifteen|twenty|thirty|forty|fifty|sixty|hundred|thousand|\d+)\b",
 14: r"\b(am|is|are|'m|'s|'re)\b.*\w+ing\b", 16: r"\b(will|won't|'ll)\b", 17: r"\bgoing to\b",
 18: r"\b(will|won't|'ll|going to)\b", 21: r"\bthere\b.*\b(is|are|isn't|aren't)\b|^(is|are)\b",
 22: r"(er\b|more |most |est\b|less |better|worse|further)", 23: r"\b(must|mustn't|have to|has to|had to|don't have to|doesn't have to|need)\b",
 24: r"\b(should|shouldn't)\b", 25: r"ly\b|\bwell\b|\bhard\b|\bfast\b|\blate\b",
 31: r"\b(was|were)\b.*\w+ing\b", 32: r"\b(has|have)\b.*\bbeen\b.*\w+ing\b", 33: r"\b(was|were)\b.*ing\b|ed\b|\b\w+\b",
 35: r"\bhad\b", 36: r"\bwill\b.*\bbe\b.*\w+ing\b|\b'll\b.*\bbe\b.*ing\b", 38: r"\w+", 39: r"\b(will|won't|'ll)\b|\w+s\b|^(if|unless)",
 40: r"\b(would|wouldn't|could|'d|were|was|had|didn't|\w+ed)\b", 41: r"\b(am|is|are|was|were|been|being|be|get|got)\b.*\w+(ed|en|t|wn|ne|d)\b",
 43: r"\b(must|might|may|could|can't|couldn't|should|can)\b", 44: r"\b(who|which|whose|that|where|when|whom|-)\b",
 45: r"\b(would|had|was|were|could|\w+ed)\b", 46: r"\b(used to|would|used|didn't use to)\b", 47: r"\w+ing\b|\bto \w+",
 48: r"\?$", 49: r"\b(so|neither|nor)\b", 50: r"\bhad been\b", 51: r"\bwill have\b|\b'll have\b", 52: r"\bwill have been\b",
 53: r"\b(would|wouldn't|could|might)\b.*\bhave\b", 54: r"\b(would|wouldn't|could|might|had)\b", 55: r"\b(to be|to have been|being|been)\b",
 56: r"\w+ing\b|\bto \w+|\bthat\b|\bhad\b|\bwould\b", 57: r"\b(should|shouldn't|could|couldn't|must|might|needn't)\b.*\bhave\b|\bhave\b",
 58: r"\b(do|does|did|had|have|has|is|are|was|were|can|could|will|would|should)\b", 59: r"\w+(ed|en|t|d)\b|\bto \w+|\w+ing\b",
 60: r"\b(would|had|were|was|could|\w+ed)\b", 61: r"\b(who|which|whose|that|where|when|whom|-)\b", 62: r"^(a|an|the|-)\b|\bthe\b|\ba\b|\ban\b",
 64: r"\b(would|was going to|were going to)\b", 65: r"\b(is|are|am|has|have|'s|'re|'m)\b|\w+s\b", 66: r"\b(was|were|had|'d)\b|\w+ed\b", 67: r"\b(will|'ll|going to|won't)\b",
}
NO_INTRO = {27, 68, 74, 75}


def s(v):
    return "" if v is None else str(v).strip()


def main():
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    types = {int(r[0]): s(r[1]) for r in wb["ALIAS - exercise_types"].iter_rows(min_row=2, values_only=True)
             if r[0] is not None and s(r[0]).isdigit()}
    ex = {int(r[0]): (int(r[2]), int(r[3])) for r in wb["exercises"].iter_rows(min_row=2, values_only=True)
          if r[0] is not None and s(r[0]).isdigit()}
    hits = 0
    for r in wb["sentence_translations"].iter_rows(min_row=2, values_only=True):
        if r[0] is None or not s(r[0]).isdigit() or s(r[2]) != "en" or not s(r[4]):
            continue
        eid = int(r[1]); mid, t = ex[eid]
        intro, ca = s(r[3]), s(r[4])
        pat = PAT.get(t)
        if pat and not re.search(pat, (intro + " " + ca if t in (14, 21, 32, 36, 41) else ca).lower()):
            hits += 1
            print(f"{mid} {eid} type {t} {types.get(t,'')[:26]:<26} | {intro[:60]} | {ca}")
    print(f"{hits} rows to eyeball")


if __name__ == "__main__":
    main()
