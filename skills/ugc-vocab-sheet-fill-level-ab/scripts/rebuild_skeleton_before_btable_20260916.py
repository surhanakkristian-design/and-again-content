#!/usr/bin/env python3
"""
rebuild_skeleton.py — rebuild an EMPTY part's `exercises` and `sentence_translations` rows to
BRIEF §0r / §0s. The selection itself lives in `exercise_selection.py`; this script only writes.

  python3 rebuild_skeleton.py --workbook part.xlsx --exercise-id-start N            # dry run
  ... --apply --selection-csv out.csv                                                # writes

Main-block exercises only (grammar + the Label type 27/68). The Simple Explanation (74/75) is
added afterwards by `add_simple_explanation.py` with an explicit id from the SE block, as before.

Refuses a workbook in which ANY sentence_translations text cell is filled — rebuilding a part
that carries content would destroy it. Ids: exercises dense from --exercise-id-start in the
order the media already appear in the sheet (the build's alphabetical concept order);
sentence_translations ids follow `(exercise_id - 1) * 9 + k`, nine rows per exercise.
Per media the rows are ordered by type id, the Label last. options_count = 3 everywhere
(provisional, as `build_parts.py` writes it; §0h.3's cut to 2 happens after the English pass).
"""
import argparse, csv, os, sys
from collections import Counter
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import exercise_selection as X

LANGS = ["sk", "en", "de", "cz", "fr", "es", "ua", "tr", "hu"]
LABEL = {"A": 27, "B": 68}


def s(v):
    return "" if v is None else str(v).strip()


def main(a):
    wb = openpyxl.load_workbook(a.workbook)
    aw = {int(r[0]): r for r in wb["All Words"].iter_rows(min_row=2, values_only=True) if s(r[0]).isdigit()}
    md = {int(r[0]): r for r in wb["media"].iter_rows(min_row=2, values_only=True) if s(r[0]).isdigit()}
    con = {int(r[0]): (s(r[1]), s(r[2]).lower()) for r in wb["word_concepts"].iter_rows(min_row=2, values_only=True) if s(r[0]).isdigit()}
    cm = {int(r[2]): int(r[1]) for r in wb["concept_media"].iter_rows(min_row=2, values_only=True) if s(r[0]).isdigit()}
    levels = {s(r[3]).upper() for r in aw.values()}
    if len(levels) != 1:
        sys.exit(f"FATAL: mixed or missing level {levels} — one level per part")
    level = levels.pop()

    ex_ws, st_ws = wb["exercises"], wb["sentence_translations"]
    filled = 0
    for r in st_ws.iter_rows(min_row=2, values_only=True):
        if s(r[0]).isdigit() and any(s(v) for v in r[3:10]):
            filled += 1
    if filled:
        sys.exit(f"REFUSED: {filled} sentence_translations rows carry text — this part is not empty")

    # media order as the sheet already holds it (first appearance), then any media without rows
    order = []
    for r in ex_ws.iter_rows(min_row=3, values_only=True):
        if s(r[0]).isdigit() and int(r[2]) not in order:
            order.append(int(r[2]))
    order += [m for m in sorted(md) if m not in order]
    if set(order) != set(md) or set(md) != set(aw):
        sys.exit("FATAL: media / All Words / exercises disagree on the media set")

    media = [(m, s(md[m][5]), con[cm[m]][1]) for m in md]
    assign, counts, skips = X.select_part(media, level)
    other = sorted(m for m, k, p in media if X.is_other_pos(p))

    rows, eid = [], a.exercise_id_start
    for m in order:
        for t in sorted(assign[m]) + [LABEL[level]]:
            rows.append((eid, cm[m], m, t, 3))
            eid += 1
    first, last = a.exercise_id_start, eid - 1

    kinds = Counter(s(md[m][5]) for m in md)
    per_kind = Counter()
    for m in md:
        per_kind[s(md[m][5])] += len(assign[m]) + 1
    print(f"{os.path.basename(a.workbook)} — level {level}, {len(md)} media {dict(kinds)}")
    print(f"  main-block exercises {len(rows)}  ids {first}-{last}  (+{len(md)} Simple Explanation later)")
    print(f"  by media kind (grammar + Label): {dict(per_kind)}")
    print(f"  per grammar type: " + " ".join(f"{t}:{counts.get(t, 0)}" for t in X.GRAMMAR_TYPES[level]))
    zero = [t for t in X.GRAMMAR_TYPES[level] if not counts.get(t)]
    print(f"  types at zero: {zero or 'none'}")
    print(f"  words outside noun/verb/adjective/adverb (group 2+4): {len(other)} {other[:10]}")
    print(f"  images whose round-robin type was skipped for part of speech: {len(skips)}")
    for m, sk in sorted(skips.items()):
        print(f"    media {m} ({con[cm[m]][1]} '{con[cm[m]][0]}'): skipped {sk} -> {assign[m]}")

    if a.selection_csv:
        with open(a.selection_csv, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["media_id", "media_type", "concept_id", "word", "part_of_speech", "grammar_types", "image_skipped"])
            for m in sorted(md):
                w.writerow([m, s(md[m][5]), cm[m], con[cm[m]][0], con[cm[m]][1],
                            " ".join(map(str, assign[m])), " ".join(map(str, skips.get(m, [])))])
    if not a.apply:
        print("\ndry run — add --apply to write")
        return

    old_ex = sum(1 for r in ex_ws.iter_rows(min_row=3, values_only=True) if s(r[0]).isdigit())
    old_st = sum(1 for r in st_ws.iter_rows(min_row=2, values_only=True) if s(r[0]).isdigit())
    ex_ws.delete_rows(3, ex_ws.max_row)
    st_ws.delete_rows(2, st_ws.max_row)
    ri, si = 3, 2
    for (e, c, m, t, oc) in rows:
        for col, v in enumerate((e, c, m, t, oc), start=1):
            ex_ws.cell(ri, col, v)
        ri += 1
        for k, lang in enumerate(LANGS):
            st_ws.cell(si, 1, (e - 1) * 9 + k + 1)
            st_ws.cell(si, 2, e)
            st_ws.cell(si, 3, lang)
            si += 1
    wb.save(a.workbook)
    print(f"\nwritten: exercises {old_ex} -> {len(rows)}, sentence_translations {old_st} -> {len(rows) * 9}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--exercise-id-start", type=int, required=True)
    ap.add_argument("--selection-csv", default=None)
    ap.add_argument("--apply", action="store_true")
    main(ap.parse_args())
