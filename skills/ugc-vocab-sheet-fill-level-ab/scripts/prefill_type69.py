#!/usr/bin/env python3
"""prefill_type69.py — writes the part of every type-69 (Meaning of the word) English row that is never authored.

    python3 prefill_type69.py <workbook.xlsx> <out_dir>

THIS is where the type-69 stem is written (PROMPT 2026-09-16/03, task 5). For each type-69 exercise it writes
<out_dir>/t69_<media_id>.txt holding

    E <exercise_id>
    en|“<English headword>” means...|<meaning of the word, All Words column F, verbatim>|<distractor_1>|<distractor_2>

with the two distractor fields EMPTY for the writer to fill. The stem follows the English HEADWORD, not
word_concepts.word stripped: a verb reads “to adore” means... (Kristian 16.9.2026: "Ja by som dal 'to adore' nech je
jasné, že sa jedná o sloveso"), every other part of speech — the -ing keywords included, which are nouns — the bare
word. The headword rule is apply_en.en_headword(), the same function the gate checks with. The meaning is copied
from the cell character for character; the file is read back and compared before the script exits.
Also writes <out_dir>/W_en_headwords.txt (the W blocks of the English headwords, same rule). Read-only on the workbook."""
import os, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from apply_en import en_headword, s

wbp, out = sys.argv[1:3]
os.makedirs(out, exist_ok=True)
wb = openpyxl.load_workbook(wbp, read_only=True)
con = {r[0]: (r[1], r[2]) for r in wb["word_concepts"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)}
raw_meaning = {r[0]: r[5] for r in wb["All Words"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)}
rows, problems = [], []
for r in wb["exercises"].iter_rows(min_row=2, values_only=True):
    if isinstance(r[0], int) and r[3] == 69:
        eid, cid, mid = r[0], r[1], r[2]
        word, pos = con[cid]; meaning = raw_meaning.get(mid)
        if meaning is None or s(meaning) != meaning: problems.append(f"media {mid}: meaning cell {meaning!r} is empty or carries edge whitespace")
        elif "|" in meaning or "\n" in meaning: problems.append(f"media {mid}: meaning carries a pipe or newline")
        elif len(meaning) > 50: problems.append(f"media {mid}: meaning {len(meaning)} chars > 50")
        rows.append((mid, eid, f"“{en_headword(word, pos)}” means...", meaning, word, pos))
if problems:
    print("STOP —", len(problems), "problem(s):"); [print("  " + p) for p in problems]; sys.exit(1)
for mid, eid, stem, meaning, word, pos in rows:
    open(os.path.join(out, f"t69_{mid}.txt"), "w", encoding="utf-8").write(f"E {eid}\nen|{stem}|{meaning}||\n")
# read back: every answer field must equal the cell exactly
bad = 0
for mid, eid, stem, meaning, word, pos in rows:
    line = open(os.path.join(out, f"t69_{mid}.txt"), encoding="utf-8").read().split("\n")[1].split("|")
    bad += (line[1] != stem or line[2] != meaning)
heads = sorted({cid: en_headword(*con[cid]) for cid in con}.items())
open(os.path.join(out, "W_en_headwords.txt"), "w", encoding="utf-8").write(
    "# W blocks — English headwords: word_concepts.word verbatim; verbs take 'to ' (brief §0t); -ing keywords are nouns and stay bare.\n\n"
    + "".join(f"W {cid}\nen|{h}\n" for cid, h in heads))
import collections
print(f"type-69 rows: {len(rows)} | files: {len(rows)} | read-back mismatches: {bad} | stems by POS: "
      f"{dict(collections.Counter(p for *_, p in rows))} | verb stems with 'to ': {sum(1 for r in rows if r[5]=='verb' and r[2].startswith(chr(0x201c)+'to '))}"
      f" | meanings containing the word: {[(m, w) for m, e, st, mean, w, p in rows if w.lower() in mean.lower()]} | headwords: {len(heads)}")
sys.exit(1 if bad else 0)
