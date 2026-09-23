import openpyxl, sys, glob, os, json
roots=[os.path.expanduser(p) for p in ["~/Documents/Claude","~/Meine Ablage/And Again/Excels","~/Meine Ablage/Excels"]]
files=[]
for r in roots:
    for f in glob.glob(r+"/**/*.xlsx",recursive=True):
        if "/~$" in f or "/logs/" in f: continue
        files.append(f)
out=[]
for f in sorted(set(files)):
    try:
        wb=openpyxl.load_workbook(f,read_only=True)
    except Exception as e:
        print("ERR",f,e);continue
    for sn in wb.sheetnames:
        if sn.strip().lower()!="all words": continue
        ws=wb[sn]
        for i,row in enumerate(ws.iter_rows(min_row=1,max_row=5,values_only=True)):
            h=[str(c).strip() if c is not None else "" for c in row]
            if any(x.lower()=="media_id" for x in h) or any(x.lower()=="voiceover" for x in h):
                print(f, "| hdr row",i+1,"|",[x for x in h if x][:25]);break
        else: print(f,"| no hdr", )
    wb.close()
