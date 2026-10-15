import openpyxl, json, os
OUT=os.path.expanduser("~/Projects/and-again-content/runs/a102_20261015/excel/")
D="/Users/kristiansurhanak/Meine Ablage/And Again/Claude Krasty/"
FILES=[D+"1.Step - Exercises/And_Again_coverage_v3.xlsx",D+"And_Again_coverage_v4.xlsx"]
exp={(c['file'],c['sheet'],c['cell']):c['new'] for c in json.load(open(OUT+"changes102.json"))}
ok=True
for p in FILES:
    a=openpyxl.load_workbook(p.replace(".xlsx","_backup_a102.xlsx")); b=openpyxl.load_workbook(p); n=os.path.basename(p)
    assert a.sheetnames==b.sheetnames
    for s in a.sheetnames:
        wa,wb=a[s],b[s]
        if (wa.max_row,wa.max_column)!=(wb.max_row,wb.max_column): ok=False; print("DIM",n,s)
        for r in range(1,max(wa.max_row,wb.max_row)+1):
            for c in range(1,max(wa.max_column,wb.max_column)+1):
                x,y=wa.cell(r,c),wb.cell(r,c); k=(n,s,y.coordinate)
                if k in exp:
                    if y.value!=exp[k]: ok=False; print("WRONG",k,y.value)
                elif x.value!=y.value: ok=False; print("DIFF",k,x.value,y.value)
                if x.style_id!=y.style_id and (x.number_format!=y.number_format or x.font!=y.font or x.fill!=y.fill or x.border!=y.border): ok=False; print("STYLE",k)
    print(n,"checked")
print("PASS" if ok else "FAIL", len(exp),"intended cells")
