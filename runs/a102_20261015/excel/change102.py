import openpyxl, shutil, json, os
D="/Users/kristiansurhanak/Meine Ablage/And Again/Claude Krasty/"
OUT=os.path.expanduser("~/Projects/and-again-content/runs/a102_20261015/excel/")
FILES={"W3":D+"1.Step - Exercises/And_Again_coverage_v3.xlsx","W4":D+"And_Again_coverage_v4.xlsx"}
OLD_IDS="3472, 4647"; NEW={"media_ids":"3472, 4647, 43","media_count":3}
changes=[]
for key,p in FILES.items():
    b=p.replace(".xlsx","_backup_a102.xlsx")
    assert not os.path.exists(b), b
    shutil.copy2(p,b); shutil.copy2(p,OUT+"backup/"+os.path.basename(b))
    wb=openpyxl.load_workbook(p); ws=wb['Target']; H={c.value:i+1 for i,c in enumerate(ws[1])}
    rows=[r for r in range(2,ws.max_row+1) if ws.cell(r,H['word']).value=="spin"]
    assert len(rows)==1, rows; r=rows[0]
    assert ws.cell(r,H['media_ids']).value==OLD_IDS and ws.cell(r,H['media_count']).value==2
    assert ws.cell(r,H['status']).value=="covered" and ws.cell(r,H['covered_by']).value=="video + picture"
    for col,v in NEW.items():
        c=ws.cell(r,H[col]); changes.append(dict(file=os.path.basename(p),sheet="Target",row=r,cell=c.coordinate,column=col,old=c.value,new=v)); c.value=v
    wb.save(p)
json.dump(changes,open(OUT+"changes102.json","w"),indent=1,ensure_ascii=False)
for c in changes: print(c)
