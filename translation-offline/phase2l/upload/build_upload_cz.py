#!/usr/bin/env python3
"""Phase 2L Part D - the Czech final upload (0 model calls), mirroring phase2k/build_upload.py for Slovak: rows from
phase2h/out/annotations_cz_final.jsonl, sheet `cz`, columns = 2K COLS, structure_json = the jsonl row; 4,064 rows,
exercise_id integer, v[0] == en; `v` display-only; 2K's reference-ending scan (phase2k/build_upload.scan_endings) applied."""
import collections, importlib.util, json, os, sys
sys.dont_write_bytecode = True
UP = os.path.dirname(os.path.abspath(__file__)); HERE = os.path.dirname(UP); TOFF = os.path.dirname(HERE)
SRC = os.path.join(TOFF, 'phase2h', 'out', 'annotations_cz_final.jsonl')
OUT_JL, OUT_X = os.path.join(UP, 'annotations_cz_final.jsonl'), os.path.join(UP, 'upload_cz_final.xlsx')
spec = importlib.util.spec_from_file_location('build_upload_2k', os.path.join(TOFF, 'phase2k', 'build_upload.py'))
BU = importlib.util.module_from_spec(spec); spec.loader.exec_module(BU)
jl = BU.jl
rows = jl(SRC)
# 2I Part 0 correction, mirrored (phase2i/upload/UPLOAD_README.md): `v = [en]` where v is empty (CZ: 50 rows)
V_EMPTY = []
for r in rows:
    if r.get('v') == [] and isinstance(r.get('en'), str) and r['en']:
        r['v'] = [r['en']]; V_EMPTY.append(r['exercise_id'])
ids = [r.get('exercise_id') for r in rows]
pre = {'rows': len(rows), 'unique_ids': len(set(ids)), 'int_ids': all(type(i) is int for i in ids),
       'v0_eq_en_violations': [r.get('exercise_id') for r in rows if not (isinstance(r.get('v'), list) and r['v'] and r['v'][0] == r.get('en'))],
       'language_code': dict(collections.Counter(r.get('language_code') for r in rows))}
ok_pre = pre['rows'] == 4064 == pre['unique_ids'] and pre['int_ids'] and not pre['v0_eq_en_violations'] and pre['language_code'] == {'cz': 4064}
chk = {'source': 'phase2h/out/annotations_cz_final.jsonl', 'v_empty_set_to_en': V_EMPTY, 'pre': pre, 'pre_ok': ok_pre}
_j2 = [json.loads(l) for l in open(os.path.join(TOFF, 'phase2j', 'upload', 'annotations_cz_fixed.jsonl'), encoding='utf-8') if l.strip()]
chk['rows_identical_to_phase2i_2j_cz_fixed'] = sum(1 for a, b in zip(rows, _j2) if a == b)
if ok_pre:
    with open(OUT_JL, 'w', encoding='utf-8') as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + '\n')
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'cz'; ws.append(list(BU.COLS))
    for r in rows:
        ws.append([r['exercise_id'], r['language_code'], r['level'], r['src'], r['en'], json.dumps(r, ensure_ascii=False)])
    wb.save(OUT_X)
    # --- validate the written file
    wb2 = openpyxl.load_workbook(OUT_X, read_only=True); ws2 = wb2['cz']
    it = ws2.iter_rows(values_only=True); hdr = list(next(it)); data = list(it)
    post = {'sheets': wb2.sheetnames, 'header': hdr, 'header_ok': hdr == list(BU.COLS), 'data_rows': len(data),
            'exercise_id_cell_types': dict(collections.Counter(type(d[0]).__name__ for d in data)),
            'unique_ids': len({d[0] for d in data}),
            'v0_eq_en_col': sum(1 for d in data if json.loads(d[5])['v'][0] == d[4]),
            'structure_id_eq_col': sum(1 for d in data if json.loads(d[5])['exercise_id'] == d[0]),
            'rows_equal_source': sum(1 for d, r in zip(data, rows) if json.loads(d[5]) == r),
            'jsonl_equal_source': jl(OUT_JL) == rows}
    try:
        import pandas as pd
        df = pd.read_excel(OUT_X, sheet_name='cz')
        post['pandas_rows'] = len(df); post['pandas_exercise_id_dtype'] = str(df['exercise_id'].dtype)
    except Exception as e:
        post['pandas'] = 'unavailable: %s' % e
    chk['post'] = post
    chk['post_ok'] = (post['header_ok'] and post['data_rows'] == 4064 == post['unique_ids'] == post['v0_eq_en_col']
                      == post['structure_id_eq_col'] == post['rows_equal_source'] and post['exercise_id_cell_types'] == {'int': 4064}
                      and post['jsonl_equal_source'] and post.get('pandas_exercise_id_dtype', 'int64') == 'int64')
    hits = BU.scan_endings(OUT_JL, 'cz')
    rev = {(x['exercise_id'], x['ref']) for x in json.load(open(os.path.join(TOFF, 'phase2k', 'analysis', 'ref_endings_reviewed.json'), encoding='utf-8')) if x['lang'] == 'cz'}
    for h in hits:
        h['reviewed_2k'] = (h['exercise_id'], h['ref']) in rev
    k2 = [h for h in json.load(open(os.path.join(TOFF, 'phase2k', 'analysis', 'ref_endings.json'), encoding='utf-8')) if h['lang'] == 'cz']
    key = lambda h: (h['exercise_id'], h['ref_index'], h['ref'])
    json.dump(hits, open(os.path.join(HERE, 'partD', 'ref_endings_cz.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    chk['ref_endings'] = {'broad': len(hits), 'strict': sum(h['strict'] for h in hits),
                          'strict_unreviewed': [[h['exercise_id'], h['ref_index'], h['last'], h['ref']] for h in hits if h['strict'] and not h['reviewed_2k']],
                          'by_last': dict(collections.Counter(h['last'] for h in hits)),
                          'vs_2k_cz_scan (phase2i/upload file)': {'2k_hits': len(k2), 'same': len({key(h) for h in hits} & {key(h) for h in k2}),
                                                                   'new_here': len({key(h) for h in hits} - {key(h) for h in k2}),
                                                                   'gone_here': len({key(h) for h in k2} - {key(h) for h in hits})}}
    j2 = {r['exercise_id']: r for r in jl(os.path.join(TOFF, 'phase2j', 'upload', 'annotations_cz_fixed.jsonl'))}
    b3 = [d for d in jl(os.path.join(TOFF, 'phase2j', 'upload', 'B3_diff.jsonl')) if d.get('lang') == 'cz']
    chk['vs_phase2j_cz_fixed'] = {'rows_2j': len(j2), 'same_ids': set(j2) == set(ids),
                                  'v_differs': sum(1 for r in rows if r['exercise_id'] in j2 and j2[r['exercise_id']].get('v') != r['v']),
                                  'b3_cz_entries': len(b3), 'b3_cz_actions': dict(collections.Counter(d.get('action') for d in b3))}
json.dump(chk, open(os.path.join(UP, 'upload_cz_check.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
re_ = chk.get('ref_endings', {})
md = ['# Phase 2L - Czech upload file check (UPLOAD_CZ_CHECK)', '',
      'File `phase2l/upload/upload_cz_final.xlsx` (sheet `cz`, columns %s) + `annotations_cz_final.jsonl`; built by `build_upload_cz.py` '
      'from `phase2h/out/annotations_cz_final.jsonl`, mirroring 2K `build_upload.py`.' % ', '.join(BU.COLS), '',
      '- 2I Part 0 correction mirrored: `v = [en]` on %d rows with empty `v` (%s); rows identical to phase2i/phase2j `annotations_cz_fixed.jsonl`: %d / 4,064' % (len(V_EMPTY), ', '.join(map(str, V_EMPTY)), chk['rows_identical_to_phase2i_2j_cz_fixed']),
      '- source checks: %s -> %s' % (json.dumps(pre, ensure_ascii=False), 'PASS' if ok_pre else 'FAIL (file NOT written)')]
if ok_pre:
    md += ['- written-file checks: %s -> %s' % (json.dumps(chk['post'], ensure_ascii=False), 'PASS' if chk['post_ok'] else 'FAIL'),
           '- `v` is display-only: the checker is SOURCE-ONLY and never grades against `en` / `v`.',
           '- 2K reference-ending scan on Czech: %d broad hits, %d strict (article / non-particle preposition); strict not reviewed in 2K: %s; '
           'by last word %s; vs 2K\'s CZ scan: %s. Detail `partD/ref_endings_cz.json`.' % (
               re_['broad'], re_['strict'], json.dumps(re_['strict_unreviewed'], ensure_ascii=False), re_['by_last'],
               re_['vs_2k_cz_scan (phase2i/upload file)']),
           '- vs `phase2j/upload/annotations_cz_fixed.jsonl`: %s' % json.dumps(chk['vs_phase2j_cz_fixed'])]
open(os.path.join(UP, 'UPLOAD_CZ_CHECK.md'), 'w', encoding='utf-8').write('\n'.join(md) + '\n')
if ok_pre:
    open(os.path.join(UP, 'UPLOAD_README.md'), 'w', encoding='utf-8').write(
        '# Phase 2L - final Czech upload\n\n`upload_cz_final.xlsx` (sheet `cz`) + `annotations_cz_final.jsonl`, 4,064 rows, integer '
        '`exercise_id`, `v[0] == en` on every row. **`v` is display-only and does not grade**: the Czech checker (stack_2l_cz) '
        'is SOURCE-ONLY. Checks: `UPLOAD_CZ_CHECK.md`.\n')
print('\n'.join(md))
