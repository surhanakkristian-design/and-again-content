#!/usr/bin/env python3
"""Phase 2J S5 - B2 summary, B3 reference corrections (flagged only), B4 closed-set re-score set + scoring, PART_B.md.
Usage (absolute paths):  python3 -B b3.py build | selftest | score
0 model calls in this file. Writes only under phase2j/."""
import ast, collections, hashlib, json, math, os, re, shutil, sys

P2J = '/Users/kristiansurhanak/Projects/and-again-content/translation-offline/phase2j'
TOFF = os.path.dirname(P2J); P2I = TOFF + '/phase2i'
PB = P2J + '/partB'; AUD = PB + '/audit'; UP = P2J + '/upload'; B4 = PB + '/b4'
sys.path.insert(0, P2J + '/partB')
import lever3  # the copy in phase2j/partB (asserted equal to phase1p/lever3.py in build)

def guard(p):
    p = os.path.abspath(p)
    assert p.startswith(P2J + '/'), 'write outside phase2j refused: ' + p
    return p
def jl(p): return [json.loads(l) for l in open(p, encoding='utf-8') if l.strip()]
def wj(p, o): json.dump(o, open(guard(p), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
def wjl(p, rows):
    with open(guard(p), 'w', encoding='utf-8') as f:
        for r in rows: f.write(json.dumps(r, ensure_ascii=False) + '\n')
def sha(p): return hashlib.sha256(open(p, 'rb').read()).hexdigest()

# ---------------------------------------------------------------- transforms
TF = lever3.time_frame
MALE = {'he': 'she', 'him': 'her', 'his': 'her', 'himself': 'herself'}
FEM = {'she': 'he', 'herself': 'himself', 'hers': 'his'}          # 'her' resolved by context
HER_OBJ_NEXT = set('to for with at in on and or but up down out off back away again that so a an the about from by into over too '
                   'now today yesterday tomorrow all very more because when if before after while as his her my your our their its '
                   'this these those some any every no than here there home'.split())
GWORD = re.compile(r"\b(he|him|his|himself|she|her|hers|herself)\b", re.I)
def case_like(src, new):
    return new[:1].upper() + new[1:] if src[:1].isupper() else new
def fam(w):
    w = w.lower()
    return 'm' if w in MALE else 'f' if (w in FEM or w == 'her') else None
def gender_swap(s):
    """swap every he/she-family pronoun; only for sentences whose gendered pronouns are ALL one gender (else None)."""
    fams = {fam(m.group(1)) for m in GWORD.finditer(s)}
    if len(fams) != 1: return None
    def rep(m):
        w = m.group(1); lw = w.lower()
        if lw == 'her':
            nxt = re.match(r"\s*([A-Za-z']+)", s[m.end():])
            new = 'him' if (not nxt or nxt.group(1).lower() in HER_OBJ_NEXT or s[m.end():m.end()+1] in ',.!?;:') else 'his'
        else:
            new = MALE.get(lw) or FEM[lw]
        return case_like(w, new)
    return GWORD.sub(rep, s)
def span_re(span): return re.compile(r"(?<![A-Za-z'])" + re.escape(span) + r"(?![A-Za-z'])", re.I)
ART_BAD = [(re.compile(r"\b([Aa])n ([bcdfgjklmnpqrstvwxyz])"), r"\1 \2"), (re.compile(r"\b([Aa]) ([aeiou])"), r"\1n \2")]
def clean(s):
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'\s+([,.!?;:])', r'\1', s)
    s = re.sub(r'([,;:])(\s*[,;:])+', r'\1', s)
    s = re.sub(r'^[,;:\s]+', '', s)
    return s[:1].upper() + s[1:] if s else s
def remove_span(ref, span):
    """A class: delete the added span (exactly one occurrence). Returns (new, reason_if_refused)."""
    if not span or not span.strip(): return None, 'no_span'
    ms = list(span_re(span.strip()).finditer(ref))
    if len(ms) != 1: return None, 'span_found_%d' % len(ms)
    if len(span.split()) * 2 > len(ref.split()): return None, 'span_too_large'
    m = ms[0]; new = ref[:m.start()] + ' ' + ref[m.end():]
    had = any(p.search(ref) for p, _ in ART_BAD)
    new = clean(new)
    if not had:
        for p, r in ART_BAD: new = p.sub(r, new)
    if TF(new) != TF(ref): return None, 'tense_changed'
    return new, None
def replace_span(ref, span, alt):
    """N class: the broader word as a further variant."""
    if not span or not alt or not span.strip() or not alt.strip(): return None, 'no_span_or_alt'
    ms = list(span_re(span.strip()).finditer(ref))
    if len(ms) != 1: return None, 'span_found_%d' % len(ms)
    m = ms[0]
    return clean(ref[:m.start()] + case_like(m.group(0), alt.strip()) + ref[m.end():]), None
def w_is_gender(span, alt):
    return bool(span and alt) and fam(span.strip()) is not None and fam(alt.strip()) is not None \
        and fam(span.strip()) != fam(alt.strip()) and len(span.split()) == 1 and len(alt.split()) == 1

def correct_row(v, flags, b1_contra=frozenset()):
    """v = stored references (v[0] == en); flags = {ref_index: audit row}. Returns (new_v, actions).
    Faithful / unflagged references are never touched. v[0] changes only if v[0] itself is flagged (A, or W contradiction)."""
    v = list(v); main_tf = TF(v[0]); acts = []; adds = []
    for i in sorted(flags):
        f = flags[i]; c = f['c']; ref = v[i]
        base = {'ref_index': i, 'class': c, 'before': ref, 'span': f.get('span'), 'alt': f.get('alt')}
        if c == 'A':
            new, why = remove_span(ref, f.get('span'))
            if new is None: acts.append(dict(base, action='listed', reason=why)); continue
            v[i] = new; acts.append(dict(base, action='replaced', after=new))
        elif c == 'N':
            new, why = replace_span(ref, f.get('span'), f.get('alt'))
            if new is None: acts.append(dict(base, action='listed', reason=why)); continue
            adds.append((new, base))
        elif c == 'W':
            if not w_is_gender(f.get('span'), f.get('alt')):
                acts.append(dict(base, action='listed', reason='W_not_a_gender_pair')); continue
            new = gender_swap(ref)
            if new is None: acts.append(dict(base, action='listed', reason='W_mixed_genders')); continue
            if i in b1_contra:   # source fixes the gender (B1 contradiction agrees) -> the flagged ref is replaced
                if TF(new) != TF(ref): acts.append(dict(base, action='listed', reason='tense_changed')); continue
                v[i] = new; acts.append(dict(base, action='replaced', after=new, reason='B1_gender_contradiction'))
            else:                # gender open -> add the other gender
                adds.append((new, base))
        else:
            acts.append(dict(base, action='listed', reason='class_O_not_corrected'))
    main_tf = TF(v[0])
    for new, base in adds:
        if TF(new) != main_tf:
            acts.append(dict(base, action='listed', reason='tense_filter_%s_vs_%s' % (TF(new), main_tf), after=new)); continue
        if new in v:
            acts.append(dict(base, action='listed', reason='duplicate_variant', after=new)); continue
        v.append(new); acts.append(dict(base, action='added', after=new, ref_index_new=len(v) - 1))
    # dedup (an A removal may equal another variant) -> the duplicate is removed, never v[0]
    out = []
    for i, x in enumerate(v):
        if x in out: acts.append({'ref_index': i, 'class': 'dedup', 'action': 'removed', 'before': x})
        else: out.append(x)
    return out, acts

# ---------------------------------------------------------------- selftest (called from test_2j T17)
def selftest():
    # tense filter (lever3 copy)
    assert TF('She will hit him.') == 'future' and TF('She hugged him yesterday.') == 'past' and TF('She hits him.') == 'present'
    assert sha(P2J + '/partB/lever3.py') == sha(TOFF + '/phase1p/lever3.py')
    # A: remove added content (+ article repair), v[0] may change, en==v[0] kept by caller
    nv, a = correct_row(['There are four different charts on the bright screen.'], {0: {'c': 'A', 'span': 'bright'}})
    assert nv == ['There are four different charts on the screen.'] and a[0]['action'] == 'replaced', (nv, a)
    nv, a = correct_row(['She saw an incredible owl.'], {0: {'c': 'A', 'span': 'incredible'}})
    assert nv == ['She saw an owl.'], nv
    nv, a = correct_row(['He bought an old car.'], {0: {'c': 'A', 'span': 'old'}})
    assert nv == ['He bought a car.'], nv
    nv, a = correct_row(['Anyway, he left.'], {0: {'c': 'A', 'span': 'Anyway'}})
    assert nv == ['He left.'], nv
    nv, a = correct_row(['He left the the room.'], {0: {'c': 'A', 'span': 'the'}})
    assert nv == ['He left the the room.'] and a[0]['action'] == 'listed', a
    # N: broader word ADDED, original kept
    nv, a = correct_row(['You should keep a charger in your bag.'], {0: {'c': 'N', 'span': 'keep', 'alt': 'carry'}})
    assert nv == ['You should keep a charger in your bag.', 'You should carry a charger in your bag.'], nv
    # W gender open -> other gender added (her -> his/him by context)
    nv, a = correct_row(['By midnight, he will have cried more tears than his sister.'], {0: {'c': 'W', 'span': 'he', 'alt': 'she'}})
    assert nv[1] == 'By midnight, she will have cried more tears than her sister.', nv
    nv, a = correct_row(['She lifts her cup and thanks her.'], {0: {'c': 'W', 'span': 'She', 'alt': 'He'}})
    assert nv[1] == 'He lifts his cup and thanks him.', nv
    # W with B1 contradiction -> replaced in place (v[0] flagged)
    nv, a = correct_row(['He said it was late.'], {0: {'c': 'W', 'span': 'He', 'alt': 'She'}}, frozenset({0}))
    assert nv == ['She said it was late.'] and a[0]['action'] == 'replaced', nv
    # W mixed genders / not a gender pair -> listed, unchanged
    nv, a = correct_row(['He gave her his jacket.'], {0: {'c': 'W', 'span': 'He', 'alt': 'She'}})
    assert nv == ['He gave her his jacket.'] and a[0]['reason'] == 'W_mixed_genders', a
    nv, a = correct_row(['You catch leaves.'], {0: {'c': 'W', 'span': 'You', 'alt': 'I'}})
    assert nv == ['You catch leaves.'] and a[0]['reason'] == 'W_not_a_gender_pair'
    # tense filter: an added variant may not change the time frame of v[0]
    nv, a = correct_row(['She hugged him yesterday.', 'She will hug him.'], {1: {'c': 'N', 'span': 'hug', 'alt': 'embrace'}})
    assert nv == ['She hugged him yesterday.', 'She will hug him.'] and a[0]['reason'].startswith('tense_filter'), (nv, a)
    nv, a = correct_row(['They walked home.'], {0: {'c': 'A', 'span': 'walked'}})
    assert nv == ['They walked home.'] and a[0]['action'] == 'listed'
    # faithful / unflagged refs untouched, O listed only
    v = ['A.', 'The dog runs fast.']
    nv, a = correct_row(v, {1: {'c': 'O', 'span': 'fast'}})
    assert nv == v and a[0]['action'] == 'listed'
    nv, a = correct_row(v, {}); assert nv == v and a == []
    # dedup after removal never removes v[0]
    nv, a = correct_row(['The dog runs.', 'The big dog runs.'], {1: {'c': 'A', 'span': 'big'}})
    assert nv == ['The dog runs.'] and a[-1]['action'] == 'removed'
    return True

# ---------------------------------------------------------------- build
def cp(x, n):
    def cdf(k, p): return sum(math.exp(math.lgamma(n+1)-math.lgamma(i+1)-math.lgamma(n-i+1)+i*math.log(p)+(n-i)*math.log(1-p)) for i in range(0, k+1))
    def bis(f):
        lo, hi = 0.0, 1.0
        for _ in range(60):
            m = (lo+hi)/2
            if f(m): lo = m
            else: hi = m
        return (lo+hi)/2
    lo = 0.0 if x == 0 else bis(lambda p: 1-cdf(x-1, p) < 0.025)
    hi = 1.0 if x == n else bis(lambda p: cdf(x, p) > 0.025)
    return [round(100*lo, 2), round(100*hi, 2)]
def rate(x, n): return {'x': x, 'n': n, 'pct': round(100.0*x/n, 2), 'cp95': cp(x, n)}
def fmt(r): return '%d/%d = %.2f %% [%.2f, %.2f]' % (r['x'], r['n'], r['pct'], r['cp95'][0], r['cp95'][1])
CLS = {'F': 'faithful', 'A': 'adds content', 'N': 'narrows', 'W': 'wrong person or gender', 'O': 'other'}

def build():
    assert selftest()
    os.makedirs(guard(UP), exist_ok=True); os.makedirs(guard(B4), exist_ok=True)
    sk_lines = open(P2I + '/upload/annotations_sk_fixed.jsonl', encoding='utf-8').read().splitlines()
    SK = [json.loads(l) for l in sk_lines]; byid = {r['exercise_id']: r for r in SK}
    assert all(r['v'][0] == r['en'] for r in SK)
    ser = None
    for cand in (dict(ensure_ascii=False), dict(ensure_ascii=True), dict(ensure_ascii=False, separators=(',', ':'))):
        if all(json.dumps(r, **cand) == l for r, l in zip(SK, sk_lines)): ser = cand; break
    assert ser is not None, 'jsonl serialisation not reproduced'
    cov = json.load(open(AUD + '/COVERAGE.json')); done = json.load(open(AUD + '/DRIVER_DONE.json'))
    audited = {int(x) for x in cov['audited_exercise_ids']}
    AR = {}
    for r in jl(AUD + '/AUDIT_RESULTS.jsonl'): AR[(int(r['exercise_id']), int(r['ref_index']))] = r
    # ---- B2 summary
    lvl = lambda e: byid[e]['level']
    per_ref = collections.defaultdict(collections.Counter); per_sent = collections.defaultdict(collections.Counter)
    sent_cls = collections.defaultdict(set); mism = []
    for (e, i), r in AR.items():
        per_ref[lvl(e)][r['c']] += 1; per_ref['all'][r['c']] += 1; sent_cls[e].add(r['c'])
        if i >= len(byid[e]['v']) or byid[e]['v'][i] != r['ref']: mism.append([e, i])
    for e, cs in sent_cls.items():
        for L in (lvl(e), 'all'):
            per_sent[L]['n'] += 1
            per_sent[L]['all_faithful' if cs == {'F'} else 'any_flag'] += 1
            for c in cs - {'F'}: per_sent[L][c] += 1
    refs_expected = sum(len(byid[e]['v']) for e in audited)
    b2 = {'coverage': {'n_audited': len(audited), 'rows_total': len(SK), 'refs_audited': len(AR), 'refs_expected_on_audited_rows': refs_expected,
                       'packets_valid': cov.get('packets_valid'), 'packets_not_audited': cov.get('packets_not_audited'), 'stop': cov.get('stop'),
                       'audited_ids_sorted_sha256': hashlib.sha256(json.dumps(sorted(audited)).encode()).hexdigest(),
                       'audited_rows_per_level': dict(collections.Counter(lvl(e) for e in audited))},
          'headless_tokens': sum(done.get('per_session_tokens', {}).values()), 'driver_spent': done.get('spent'),
          'sessions': len(done.get('per_session_tokens', {})), 'failed': done.get('failed'),
          'per_ref': {k: dict(v) for k, v in per_ref.items()}, 'per_sentence': {k: dict(v) for k, v in per_sent.items()},
          'ref_text_mismatch': mism, 'prompt_sha256': cov.get('prompt_sha256')}
    # ---- B1 agreement (SK)
    B1S = jl(PB + '/B1_flags_sk.jsonl'); B1C = jl(PB + '/B1_flags_cz.jsonl')
    agree = collections.defaultdict(collections.Counter); b1_contra = collections.defaultdict(set); b1_keys = set()
    for f in B1S:
        k = (int(f['exercise_id']), int(f['ref_index'])); b1_keys.add(k)
        if k[0] not in audited: agree[f['class']]['not_audited'] += 1; continue
        c = AR.get(k, {}).get('c', 'missing'); agree[f['class']]['B2_' + c] += 1
        if f['class'] == 'gender_contradiction' and c == 'W': b1_contra[k[0]].add(k[1])
    b2W = [k for k, r in AR.items() if r['c'] == 'W']
    agreement = {'b1_sk_flags': len(B1S), 'by_b1_class': {k: dict(v) for k, v in agree.items()},
                 'b1_audited': sum(v[x] for v in agree.values() for x in v if x != 'not_audited'),
                 'b1_audited_B2_W': sum(v['B2_W'] for v in agree.values()),
                 'b1_audited_B2_any_flag': sum(v[x] for v in agree.values() for x in v if x not in ('not_audited', 'B2_F')),
                 'b2_W_refs': len(b2W), 'b2_W_also_b1_flag': sum(1 for k in b2W if k in b1_keys)}
    # ---- B3 corrections
    diff = []; newSK = {}
    for e in sorted(audited):
        r = byid[e]; flags = {i: AR[(e, i)] for i in range(len(r['v'])) if (e, i) in AR and AR[(e, i)]['c'] != 'F'}
        flags = {i: f for i, f in flags.items() if [e, i] not in mism}
        if not flags: continue
        nv, acts = correct_row(r['v'], flags, frozenset(b1_contra.get(e, ())))
        for a in acts: diff.append(dict(a, exercise_id=e, lang='sk', level=r['level'], src=r['src']))
        if nv != r['v']:
            nr = dict(r); nr['v'] = nv; nr['en'] = nv[0]; newSK[e] = nr
    for e, i in mism: diff.append({'exercise_id': e, 'lang': 'sk', 'ref_index': i, 'class': AR[(e, i)]['c'], 'action': 'listed', 'reason': 'audit_ref_text_mismatch'})
    for f in B1C:
        diff.append({'exercise_id': int(f['exercise_id']), 'lang': 'cz', 'level': f.get('level'), 'ref_index': f['ref_index'], 'class': 'B1_' + f['class'],
                     'action': 'listed', 'reason': 'cz_B1_flag_not_deterministic_safe (reader noise, no model audit)', 'before': f.get('ref'), 'src': f.get('src')})
    v0_changed = [e for e, r in newSK.items() if r['v'][0] != byid[e]['v'][0]]
    assert all(r['v'][0] == r['en'] for r in newSK.values())
    # ---- write upload files
    out_lines = [json.dumps(newSK[r['exercise_id']], **ser) if r['exercise_id'] in newSK else l for r, l in zip(SK, sk_lines)]
    open(guard(UP + '/annotations_sk_fixed.jsonl'), 'w', encoding='utf-8').write('\n'.join(out_lines) + '\n')
    assert open(UP + '/annotations_sk_fixed.jsonl', 'rb').read().count(b'\n') == len(SK)
    import openpyxl
    wb = openpyxl.load_workbook(P2I + '/upload/upload_sk_final.xlsx'); ws = wb['sk']
    hdr = [c.value for c in ws[1]]; ci = {h: i for i, h in enumerate(hdr)}; ser_x = None
    for row in ws.iter_rows(min_row=2):
        e = int(row[ci['exercise_id']].value); src_row = byid[e]
        if ser_x is None:
            for cand in (dict(ensure_ascii=False), dict(ensure_ascii=True)):
                if json.dumps(src_row, **cand) == row[ci['structure_json']].value: ser_x = cand
            assert ser_x is not None, 'xlsx structure_json serialisation not reproduced'
        if e in newSK:
            row[ci['en']].value = newSK[e]['en']; row[ci['structure_json']].value = json.dumps(newSK[e], **ser_x)
    wb.save(guard(UP + '/upload_sk_final.xlsx'))
    wb2 = openpyxl.load_workbook(UP + '/upload_sk_final.xlsx', read_only=True); n_rt = 0
    for row in list(wb2['sk'].iter_rows(values_only=True))[1:]:
        e = int(row[ci['exercise_id']]); exp = newSK.get(e, byid[e])
        assert json.loads(row[ci['structure_json']]) == exp and row[ci['en']] == exp['en']; n_rt += 1
    assert n_rt == len(SK)
    for f in ('annotations_cz_fixed.jsonl', 'upload_cz_final.xlsx'): shutil.copyfile(P2I + '/upload/' + f, guard(UP + '/' + f))
    wjl(UP + '/B3_diff.jsonl', diff)
    cnt = collections.defaultdict(collections.Counter)
    for d in diff: cnt['%s/%s' % (d['lang'], d['class'])][d['action']] += 1
    reasons = collections.Counter('%s/%s:%s' % (d['lang'], d['class'], d.get('reason')) for d in diff if d['action'] == 'listed')
    b3 = {'rows_changed_sk': len(newSK), 'rows_changed_cz': 0, 'v0_changed_rows': v0_changed, 'counts': {k: dict(v) for k, v in cnt.items()},
          'listed_reasons': dict(reasons), 'refs_before_sk_changed_rows': sum(len(byid[e]['v']) for e in newSK),
          'refs_after_sk_changed_rows': sum(len(r['v']) for r in newSK.values()),
          'out_sha256': {f: sha(UP + '/' + f) for f in sorted(os.listdir(UP)) if not f.endswith('.md')},
          'lever3_sha256': sha(P2J + '/partB/lever3.py'), 'serialisation': {'jsonl': ser, 'xlsx': ser_x}}
    # ---- B4 set: closed 2I set with corrected references
    items = jl(P2I + '/set/items.jsonl'); ch_items = []; set_eids = sorted({int(i['exercise_id']) for i in items})
    for it in items:
        e = int(it['exercise_id'])
        if e in newSK:
            it['v'] = list(newSK[e]['v']) if isinstance(it['v'], list) else repr(newSK[e]['v'])
            ann = it.get('annotation')
            if isinstance(ann, dict): ann = dict(ann, v=list(newSK[e]['v']), en=newSK[e]['en'])
            elif isinstance(ann, str): a2 = ast.literal_eval(ann); a2['v'] = newSK[e]['v']; a2['en'] = newSK[e]['en']; ann = repr(a2)
            it['annotation'] = ann; it['en'] = newSK[e]['en']; ch_items.append(it['jid'])
    wjl(B4 + '/items_B4.jsonl', items)
    seed = []
    for p in (P2I + '/run/ledger.jsonl', P2J + '/run_S2/ledger.jsonl', P2J + '/run_S2c/ledger.jsonl'):
        if os.path.exists(p): seed += jl(p)
    wjl(B4 + '/seed_ledger.jsonl', seed)
    b4set = {'set_sentences': len(set_eids), 'set_sentences_audited': sum(1 for e in set_eids if e in audited),
             'set_sentences_corrected': sorted(e for e in set_eids if e in newSK), 'items_with_corrected_refs': ch_items,
             'seed_rows': len(seed)}
    res = {'b2': b2, 'b1_agreement': agreement, 'b3': b3, 'b4_set': b4set}
    wj(PB + '/B3_result.json', res)
    L = ['# Phase 2J S5 - corrected upload files (B3)', '',
         'Generator: `phase2j/partB/b3.py build` (0 model calls). Sources: phase2i/upload/ (unchanged) + partB/audit/AUDIT_RESULTS.jsonl (B2) + partB/B1_flags_*.jsonl (B1).',
         'Same format as phase2i/upload/ (xlsx sheet sk/cz, columns exercise_id, language_code, level, src, en, structure_json; jsonl one annotation per row, same key order and serialisation). `v[0] == en` holds on every row.', '',
         'Rules: only B2-flagged SK references on the %d audited rows are touched (B2 coverage incomplete: packets %s not audited).' % (len(audited), ', '.join(cov.get('packets_not_audited') or [])),
         '- adds content (A): the added span is removed in place (exactly one occurrence, article a/an repaired, time frame must stay the same, else listed).',
         '- narrows (N): the flagged reference stays; the span replaced by the broader word is ADDED as a further variant.',
         '- wrong person or gender (W): only he/she-family pairs on sentences with one gender only; other gender ADDED (source open), or the flagged ref REPLACED where B1 read a fixed contradicting gender; person/number W are listed, not changed.',
         '- other (O): listed, not changed. Faithful references are never touched. Added variants pass the 1P lever-3 tense filter against v[0].',
         '- Czech: 0 changes; every B1 CZ flag listed (reader noise, no model audit).', '',
         'Rows changed: SK %d, CZ 0. v[0] changed on %d rows (all flagged v[0]; en updated with it).' % (len(newSK), len(v0_changed)), '',
         '| lang/class | replaced | added | removed | listed |', '|---|---|---|---|---|']
    for k in sorted(cnt): L.append('| %s | %d | %d | %d | %d |' % (k, cnt[k]['replaced'], cnt[k]['added'], cnt[k]['removed'], cnt[k]['listed']))
    L += ['', 'Per-reference detail: `B3_diff.jsonl` (action replaced / added / removed / listed + reason). Hashes: `partB/B3_result.json` b3.out_sha256.']
    open(guard(UP + '/UPLOAD_README.md'), 'w', encoding='utf-8').write('\n'.join(L) + '\n')
    print('B2 audited', len(audited), 'refs', len(AR), 'per_ref', dict(per_ref['all']), 'tokens', b2['headless_tokens'], 'mism', len(mism))
    print('B1agree', json.dumps(agreement)[:600])
    print('B3 rows', len(newSK), 'v0', len(v0_changed), json.dumps({k: dict(v) for k, v in cnt.items()}))
    print('B3 listed', dict(reasons))
    print('B4set', json.dumps(b4set)[:400])

# ---------------------------------------------------------------- score (after run_2j)
def score():
    res = json.load(open(PB + '/B3_result.json'))
    IT = {i['jid']: i for i in jl(P2I + '/set/items.jsonl')}; lab = {j: IT[j]['judge_label'] for j in IT}
    R2I = {r['jid']: r for r in jl(P2I + '/run/results.jsonl') if r['stack'] == 'tonly'}
    A3 = {r['jid']: r for r in jl(P2J + '/run_S2c/results.jsonl')}
    NEW = {r['jid']: r for r in jl(B4 + '/run/results.jsonl')}; st = json.load(open(B4 + '/run/RUN_STATUS.json'))
    assert len(NEW) == len(A3) == len(R2I) == 900
    def sc(d):
        c = [j for j in lab if lab[j] == 'correct']; w = [j for j in lab if lab[j] == 'wrong']
        return rate(sum(bool(d[j]['accept']) for j in c), len(c)), rate(sum(bool(d[j]['accept']) for j in w), len(w))
    s2i, sa3, sb = sc(R2I), sc(A3), sc(NEW); corr = set(res['b4_set']['items_with_corrected_refs'])
    def changes(base):
        out = []
        for j in sorted(base):
            if (bool(base[j]['accept']), base[j]['layer']) != (bool(NEW[j]['accept']), NEW[j]['layer']):
                out.append({'jid': j, 'judge': lab[j], 'answer': IT[j]['answer'], 'refs_corrected': j in corr,
                            'before': [base[j]['layer'], bool(base[j]['accept'])], 'after': [NEW[j]['layer'], bool(NEW[j]['accept'])],
                            'l3_reply': NEW[j].get('l3_reply'), 'stored': bool(NEW[j].get('seeded_reply'))})
        return out
    chA3, ch2I = changes(A3), changes(R2I)
    out = {'label': 'CLOSED-SET RE-SCORE (2I set, 900 items) with A + B (corrected references), stored replies by request hash + only new L3 calls',
           '2I': {'coverage': s2i[0], 'fa': s2i[1]}, 'A3': {'coverage': sa3[0], 'fa': sa3[1]}, 'A_plus_B': {'coverage': sb[0], 'fa': sb[1]},
           'run_status': {k: st.get(k) for k in ('status', 'requests', 'needed', 'seeded_used', 'calls_made', 'counted_total', 'spend_usd', 'uncounted_attempts')},
           'changed_vs_A3': chA3, 'changed_vs_2I': ch2I}
    wj(B4 + '/B4.json', out)
    b2, ag, b3 = res['b2'], res['b1_agreement'], res['b3']
    L = ['# Phase 2J Part B (S4 B1 + S5 B2 summary, B3, B4)', '', '## B2 model audit of the Slovak references (summary)', '',
         'Coverage: %d of %d SK rows audited (%d references), packets not audited: %s (token cap stop: %s). Audited ids: partB/audit/COVERAGE.json (sha of sorted ids %s). Part D must sample only from these.'
         % (b2['coverage']['n_audited'], b2['coverage']['rows_total'], b2['coverage']['refs_audited'], ', '.join(b2['coverage']['packets_not_audited'] or []),
            (b2['coverage']['stop'] or {}).get('why'), b2['coverage']['audited_ids_sorted_sha256'][:16]),
         'Headless tokens B2: %d over %d sessions (failed: %s). Prompt sha256 %s. Audit ref text mismatches: %d.' % (b2['headless_tokens'], b2['sessions'], b2['failed'], b2['prompt_sha256'], len(b2['ref_text_mismatch'])), '',
         '### Per reference', '', '| level | faithful | adds content | narrows | wrong person/gender | other | refs |', '|---|---|---|---|---|---|---|']
    for k in ('A1', 'A2', 'B1', 'B2', 'all'):
        c = b2['per_ref'].get(k, {}); L.append('| %s | %s | %s | %s | %s | %s | %d |' % (k, *(c.get(x, 0) for x in 'FANWO'), sum(c.values())))
    L += ['', '### Per sentence (a sentence counts in every class any of its references has)', '',
          '| level | sentences | all faithful | any flag | adds content | narrows | wrong person/gender | other |', '|---|---|---|---|---|---|---|---|']
    for k in ('A1', 'A2', 'B1', 'B2', 'all'):
        c = b2['per_sentence'].get(k, {}); L.append('| %s | %d | %d | %d | %d | %d | %d | %d |' % (k, c.get('n', 0), c.get('all_faithful', 0), c.get('any_flag', 0), *(c.get(x, 0) for x in 'ANWO')))
    L += ['', '### Agreement with B1 (SK)', '',
          'B1 SK flags %d; on audited rows %d; of those B2 = wrong person/gender %d, any B2 flag %d. B2 W references %d, of which B1-flagged %d (B1 compares only rows its reader reads).'
          % (ag['b1_sk_flags'], ag['b1_audited'], ag['b1_audited_B2_W'], ag['b1_audited_B2_any_flag'], ag['b2_W_refs'], ag['b2_W_also_b1_flag']), '',
          '| B1 class | B2 classes on the same reference |', '|---|---|']
    for k, v in sorted(ag['by_b1_class'].items()): L.append('| %s | %s |' % (k, ', '.join('%s %d' % x for x in sorted(v.items()))))
    L += ['', '## B3 corrections (phase2j/upload/, see UPLOAD_README.md)', '',
          'SK rows changed %d (references on those rows %d -> %d); CZ 0. v[0] changed on %d rows (v[0] flagged; en = v[0] kept).'
          % (b3['rows_changed_sk'], b3['refs_before_sk_changed_rows'], b3['refs_after_sk_changed_rows'], len(b3['v0_changed_rows'])), '',
          '| lang/class | replaced | added | removed | listed |', '|---|---|---|---|---|']
    for k, v in sorted(b3['counts'].items()): L.append('| %s | %d | %d | %d | %d |' % (k, v.get('replaced', 0), v.get('added', 0), v.get('removed', 0), v.get('listed', 0)))
    L += ['', 'Listed (not changed) by reason: ' + '; '.join('%s %d' % x for x in sorted(b3['listed_reasons'].items())), '',
          '## B4 closed-set re-score (2I set with A + B) — %s' % out['label'], '',
          'Set: %d sentences, %d audited, %d with corrected references (%d items). Run partB/b4/run: %s.' % (res['b4_set']['set_sentences'], res['b4_set']['set_sentences_audited'],
          len(res['b4_set']['set_sentences_corrected']), len(corr), json.dumps(out['run_status'])), '',
          '| | coverage | FA |', '|---|---|---|', '| 2I | %s | %s |' % (fmt(s2i[0]), fmt(s2i[1])), '| A3 (A fix) | %s | %s |' % (fmt(sa3[0]), fmt(sa3[1])),
          '| A + B | %s | %s |' % (fmt(sb[0]), fmt(sb[1])), '', 'Every item whose verdict changed vs A3 (%d):' % len(chA3), '',
          '| jid | judge | refs corrected | A3 | A+B | L3 | answer |', '|---|---|---|---|---|---|---|']
    for c in chA3:
        L.append('| %s | %s | %s | %s %s | %s %s | %s%s | %s |' % (c['jid'], c['judge'], 'yes' if c['refs_corrected'] else 'no', c['before'][0], 'acc' if c['before'][1] else 'rej',
                 c['after'][0], 'acc' if c['after'][1] else 'rej', c['l3_reply'] or '', ' (stored)' if c['stored'] else '', c['answer'].replace('|', '/')))
    L += ['', 'Changed vs 2I: %d items (list in partB/b4/B4.json changed_vs_2I).' % len(ch2I)]
    open(guard(PB + '/PART_B.md'), 'w', encoding='utf-8').write('\n'.join(L) + '\n')
    print('B4 2I', fmt(s2i[0]), fmt(s2i[1]), '| A3', fmt(sa3[0]), fmt(sa3[1]), '| A+B', fmt(sb[0]), fmt(sb[1]), '| changed vs A3', len(chA3), 'vs 2I', len(ch2I))
    print('RUN', json.dumps(out['run_status']))

if __name__ == '__main__':
    {'build': build, 'selftest': lambda: print('selftest', selftest()), 'score': score}[sys.argv[1]]()
